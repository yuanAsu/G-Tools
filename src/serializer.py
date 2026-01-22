from openpyxl import Workbook
from openpyxl.utils import range_boundaries

def serialize_excel_to_markdown(wb: Workbook) -> str:
    """
    Serializes an Excel workbook into a token-optimized Markdown format for LLM consumption.
    Format:
    === Sheet: SheetName ===
    Row X: [val1, val2, ...]
    """
    output_lines = []

    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]

        # 1. Handle Merged Cells (In-memory fill for serialization)
        # We process a copy/view logic conceptually, but here we modify the sheet
        # because the loader gave us a disposable wb object in memory.
        merged_ranges = list(sheet.merged_cells.ranges)
        for merged_range in merged_ranges:
            min_col, min_row, max_col, max_row = range_boundaries(str(merged_range))
            top_left_value = sheet.cell(row=min_row, column=min_col).value
            sheet.unmerge_cells(str(merged_range))
            for row in range(min_row, max_row + 1):
                for col in range(min_col, max_col + 1):
                    sheet.cell(row=row, column=col).value = top_left_value

        output_lines.append(f"\n=== Sheet: {sheet_name} ===")

        # 2. Iterate Rows
        # min_row=1 to capture headers
        for i, row in enumerate(sheet.iter_rows(min_row=1, values_only=True), start=1):
            # 3. Filter Empty Rows
            # Check if row is all None or empty strings
            if all((c is None or str(c).strip() == "") for c in row):
                continue

            formatted_values = []
            for cell_val in row:
                if cell_val is None:
                    formatted_values.append("None")
                else:
                    val_str = str(cell_val).strip()
                    # 4. Smart Truncation
                    if len(val_str) > 500:
                        val_str = val_str[:100] + "...(truncated)"
                    formatted_values.append(val_str)

            # Format: Row X: [val1, val2, ...]
            # Using brackets as requested
            row_str = f"Row {i}: [{', '.join(formatted_values)}]"
            output_lines.append(row_str)

    return "\n".join(output_lines)

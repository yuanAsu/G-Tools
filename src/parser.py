import pandas as pd
import openpyxl
from openpyxl.utils import range_boundaries
from src.loader import load_complex_excel

class ExcelParser:
    def __init__(self, filepath: str):
        self.original_filepath = filepath
        self.wb = load_complex_excel(filepath)

    def get_sheet_names(self) -> list[str]:
        return self.wb.sheetnames

    def _fill_merged_cells(self, sheet) -> None:
        """
        Unmerges all cells in the sheet and fills them with the value
        from the top-left cell of the merged range.
        """
        merged_ranges = list(sheet.merged_cells.ranges)

        for merged_range in merged_ranges:
            min_col, min_row, max_col, max_row = range_boundaries(str(merged_range))
            top_left_value = sheet.cell(row=min_row, column=min_col).value

            sheet.unmerge_cells(str(merged_range))

            for row in range(min_row, max_row + 1):
                for col in range(min_col, max_col + 1):
                    sheet.cell(row=row, column=col).value = top_left_value

    def parse_sheet(self, sheet_name: str, header_rows: int = 1) -> pd.DataFrame:
        if sheet_name not in self.wb.sheetnames:
            raise ValueError(f"Sheet {sheet_name} not found.")

        sheet = self.wb[sheet_name]
        self._fill_merged_cells(sheet)

        data = list(sheet.values)
        if not data:
            return pd.DataFrame()

        if header_rows > 0:
            raw_headers = data[:header_rows]
            body = data[header_rows:]

            final_headers = []
            num_cols = len(raw_headers[0])

            for col_idx in range(num_cols):
                col_header_parts = []
                for row_idx in range(header_rows):
                    val = raw_headers[row_idx][col_idx]
                    if val and str(val).strip():
                        col_header_parts.append(str(val).strip())
                final_headers.append(" - ".join(col_header_parts))

            df = pd.DataFrame(body, columns=final_headers)
        else:
            df = pd.DataFrame(data)

        return df

    def find_anchor(self, sheet_name: str, keyword: str) -> tuple[int, int] | None:
        """
        Finds the first cell containing the keyword.
        Returns (row, col) (1-based index).
        """
        if sheet_name not in self.wb.sheetnames:
            raise ValueError(f"Sheet {sheet_name} not found.")

        sheet = self.wb[sheet_name]
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value and keyword in str(cell.value):
                    return cell.row, cell.column
        return None

    def parse_table_at_anchor(self, sheet_name: str, anchor_keyword: str, relative_header_offset: int = 1, header_rows: int = 1, max_rows: int = None) -> pd.DataFrame:
        """
        Locates a table based on an anchor keyword and extracts it.
        """
        sheet = self.wb[sheet_name]
        self._fill_merged_cells(sheet)

        anchor = self.find_anchor(sheet_name, anchor_keyword)
        if not anchor:
            raise ValueError(f"Anchor '{anchor_keyword}' not found in {sheet_name}")

        anchor_row, anchor_col = anchor
        start_row = anchor_row + relative_header_offset

        data = []
        collected_rows = 0

        # Iterate starting from the header row
        for row in sheet.iter_rows(min_row=start_row, values_only=True):
            # Stop if we hit a completely empty row (heuristics for end of table)
            if all(c is None for c in row):
                break

            data.append(row)
            collected_rows += 1

            if max_rows and collected_rows >= (max_rows + header_rows):
                break

        if not data:
            return pd.DataFrame()

        # Process Headers
        if header_rows > 0:
            raw_headers = data[:header_rows]
            body = data[header_rows:]

            final_headers = []
            # Calculate num_cols based on max length in raw_headers to avoid index errors
            num_cols = max(len(h) for h in raw_headers) if raw_headers else 0

            for col_idx in range(num_cols):
                col_header_parts = []
                for row_idx in range(header_rows):
                    if col_idx < len(raw_headers[row_idx]):
                        val = raw_headers[row_idx][col_idx]
                        if val and str(val).strip():
                            col_header_parts.append(str(val).strip())
                final_headers.append(" - ".join(col_header_parts))

            df = pd.DataFrame(body, columns=final_headers)
        else:
            df = pd.DataFrame(data)

        # Drop columns that are completely empty (often artifacts of reading whole rows)
        df.dropna(how='all', axis=1, inplace=True)
        return df

    def parse_form(self, sheet_name: str, range_ref: str = None) -> dict:
        """
        Parses a sheet or range as a form (key-value pairs).
        Heuristic: Looks for cells ending in ':' and takes the next cell as value.
        """
        if sheet_name not in self.wb.sheetnames:
            raise ValueError(f"Sheet {sheet_name} not found.")

        sheet = self.wb[sheet_name]
        self._fill_merged_cells(sheet)

        data_dict = {}

        if range_ref:
            rows_iter = sheet[range_ref]
        else:
            rows_iter = sheet.iter_rows()

        for row in rows_iter:
            cells = list(row)
            for i in range(len(cells) - 1):
                key_cell = cells[i]
                val_cell = cells[i+1]

                key_val = key_cell.value

                if isinstance(key_val, str) and key_val.strip():
                    clean_key = key_val.strip()

                    if clean_key.endswith(":"):
                        key_name = clean_key[:-1].strip()
                        data_dict[key_name] = val_cell.value

        return data_dict

    def parse_hierarchy(self, sheet_name: str, col_index_val: int = 2) -> dict:
        """
        Parses a hierarchical sheet (e.g. financial statement) into a nested dictionary.
        Assumes Col A is the index tree (indented) and Col B is the value.
        """
        if sheet_name not in self.wb.sheetnames:
            raise ValueError(f"Sheet {sheet_name} not found.")

        sheet = self.wb[sheet_name]
        data = {}
        # Stack: (indent_level, parent_dict, key_in_parent_dict)
        # Root is level -1
        stack = [(-1, data, None)]

        rows = list(sheet.iter_rows())
        if not rows:
            return {}

        # Skip header, start row 2
        for row in rows[1:]:
            cell_key = row[0]
            # Skip empty keys
            if not cell_key.value:
                continue

            key_text = str(cell_key.value).strip()

            # Value (default to Col B -> index 1)
            val = None
            if len(row) >= col_index_val:
                val = row[col_index_val-1].value

            # Determine Indent
            indent = 0
            if cell_key.alignment and cell_key.alignment.indent:
                indent = int(cell_key.alignment.indent)
            else:
                # Fallback to leading spaces
                raw = str(cell_key.value)
                indent = len(raw) - len(raw.lstrip())

            # Pop stack to find parent
            while stack and stack[-1][0] >= indent:
                stack.pop()

            if not stack:
                stack.append((-1, data, None))

            parent_indent, parent_dict, parent_key = stack[-1]

            # Resolve container
            container = parent_dict
            if parent_key is not None:
                current_val = parent_dict[parent_key]
                if not isinstance(current_val, dict):
                    # Upgrade scalar to dict
                    new_container = {"_value": current_val}
                    parent_dict[parent_key] = new_container
                    container = new_container
                else:
                    container = current_val

            # Insert current item
            if val is None:
                 container[key_text] = {}
                 stack.append((indent, container, key_text))
            else:
                 container[key_text] = val
                 stack.append((indent, container, key_text))

        return data

    def parse_all(self):
        results = {}
        for name in self.get_sheet_names():
            results[name] = self.parse_sheet(name, header_rows=1)
        return results

import sys
import openpyxl

def analyze_excel(filepath):
    try:
        wb = openpyxl.load_workbook(filepath, data_only=False) # Keep formulas if any, or use True for values
    except Exception as e:
        print(f"Error loading {filepath}: {e}")
        return

    print(f"Analysis of: {filepath}")
    print("=" * 40)

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        print(f"\nSheet: {sheet_name}")
        print(f"  Dimensions: {ws.dimensions}")
        print(f"  Max Row: {ws.max_row}, Max Col: {ws.max_column}")

        merged_ranges = ws.merged_cells.ranges
        if merged_ranges:
            print(f"  Merged Cells ({len(merged_ranges)}):")
            for rng in merged_ranges:
                print(f"    - {rng}")
        else:
            print("  Merged Cells: None")

        print("  Preview (First 5 rows):")
        for i, row in enumerate(ws.iter_rows(min_row=1, max_row=5, values_only=True)):
            # Replace None with "" for cleaner output
            formatted_row = [str(val) if val is not None else "" for val in row]
            print(f"    {i+1}: {formatted_row}")

def main():
    if len(sys.argv) < 2:
        print("Usage: python analyze_excel.py <filepath>")
        sys.exit(1)

    analyze_excel(sys.argv[1])

if __name__ == "__main__":
    main()

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
import os

def create_v2_excel():
    wb = Workbook()

    # --- Sheet 1: Mixed Content (Anchors) ---
    ws1 = wb.active
    ws1.title = "Mixed Data"

    # Table 1
    # Row 1
    ws1.cell(row=1, column=1, value="Project Info").font = Font(bold=True)

    # Row 2 Headers
    headers1 = ["ID", "Name", "Status"]
    for col, val in enumerate(headers1, 1):
        ws1.cell(row=2, column=col, value=val)

    # Row 3, 4 Data
    data1 = [
        [1, "Alpha", "Active"],
        [2, "Beta", "Pending"]
    ]
    for r_idx, row_data in enumerate(data1, 3):
        for c_idx, val in enumerate(row_data, 1):
            ws1.cell(row=r_idx, column=c_idx, value=val)

    # Row 5, 6 Empty

    # Row 7 Anchor
    ws1.cell(row=7, column=1, value="Funding Assumptions").font = Font(bold=True)

    # Row 8 Headers
    headers2 = ["Source", "Rate", "Amount"]
    for col, val in enumerate(headers2, 1):
        ws1.cell(row=8, column=col, value=val)

    # Row 9, 10 Data
    data2 = [
        ["Bank Loan", "5%", 100000],
        ["Equity", "N/A", 50000]
    ]
    for r_idx, row_data in enumerate(data2, 9):
        for c_idx, val in enumerate(row_data, 1):
            ws1.cell(row=r_idx, column=c_idx, value=val)

    # --- Sheet 2: Hierarchy (Financial Statement) ---
    ws2 = wb.create_sheet("Financial Statement")
    ws2.append(["Item", "Value"])

    items = [
        ("Assets", 0, None),
        ("Current Assets", 1, None),
        ("Cash", 2, 5000),
        ("Inventory", 2, 3000),
        ("Fixed Assets", 0, None),
        ("Equipment", 1, 10000),
        ("Buildings", 1, 50000)
    ]

    for idx, (name, indent, val) in enumerate(items, start=2):
        cell = ws2[f'A{idx}']
        cell.value = name
        cell.alignment = Alignment(indent=indent) # Logical indentation

        if val is not None:
            ws2[f'B{idx}'] = val

    # --- Sheet 3: Form Mode (Key-Value) ---
    ws3 = wb.create_sheet("Project Form")

    # Block 1
    ws3['A2'] = "Project Name:"
    ws3['B2'] = "Apollo Mission"

    ws3['A3'] = "Project Manager:"
    ws3['B3'] = "Jules"

    # Block 2 (Side by side)
    ws3['D2'] = "Total Budget:"
    ws3['E2'] = 1500000

    ws3['D3'] = "Start Date:"
    ws3['E3'] = "2024-01-01"

    # Save
    data_dir = "data"
    if not os.path.exists(data_dir):
        os.makedirs(data_dir)

    filename = os.path.join(data_dir, "complex_v2.xlsx")
    wb.save(filename)
    print(f"Created {filename}")

if __name__ == "__main__":
    create_v2_excel()

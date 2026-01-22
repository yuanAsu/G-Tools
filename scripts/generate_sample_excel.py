from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
import os

def create_complex_excel():
    wb = Workbook()

    # Sheet 1: Sales Report
    ws1 = wb.active
    ws1.title = "Sales Report"

    # Headers
    ws1['A1'] = "Region"
    ws1.merge_cells('A1:A2')

    ws1['B1'] = "Q1 2023"
    ws1.merge_cells('B1:C1')

    ws1['D1'] = "Q2 2023"
    ws1.merge_cells('D1:E1')

    ws1['B2'] = "Sales"
    ws1['C2'] = "Profit"
    ws1['D2'] = "Sales"
    ws1['E2'] = "Profit"

    # Styles
    center_align = Alignment(horizontal='center', vertical='center')
    bold_font = Font(bold=True)

    for row in ws1.iter_rows(min_row=1, max_row=2):
        for cell in row:
            cell.alignment = center_align
            cell.font = bold_font

    # Data
    ws1['A3'] = "North"
    ws1.merge_cells('A3:A4')
    ws1['B3'] = 1000
    ws1['C3'] = 200
    ws1['D3'] = 1100
    ws1['E3'] = 220

    ws1['B4'] = 1200
    ws1['C4'] = 250
    ws1['D4'] = 1300
    ws1['E4'] = 300

    ws1['A5'] = "South"
    ws1['B5'] = 800
    ws1['C5'] = 100
    ws1['D5'] = 900
    ws1['E5'] = 150

    # Sheet 2: Inventory
    ws2 = wb.create_sheet("Inventory")
    ws2.append(["Item ID", "Name", "Stock", "Category"])
    ws2.append([101, "Widget A", 50, "Hardware"])
    ws2.append([102, "Widget B", 0, "Hardware"])
    ws2.append([103, "Gadget X", None, "Electronics"])

    # Ensure data directory exists
    # Assuming script is run from project root, or we find relative path
    base_dir = os.getcwd()
    data_dir = os.path.join(base_dir, "data")
    if not os.path.exists(data_dir):
        os.makedirs(data_dir)

    filename = os.path.join(data_dir, "complex_sample.xlsx")
    wb.save(filename)
    print(f"Created {filename}")

if __name__ == "__main__":
    create_complex_excel()
